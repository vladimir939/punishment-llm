"""Assemble the figures deck CJSJ asks for, from the template they publish.

    python tools\\make_figures_pptx.py

The journal wants figures in a separate .pptx with captions, built on
CJSJ-Figures-Template.pptx. We keep the template's title slide and its slide
geometry and replace the two example slides with our three figures.
"""

from __future__ import annotations

import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "submission" / "CJSJ-Figures-Template.pptx"
OUT = ROOT / "submission" / "VyatchaninVladimir_figures.pptx"
FIGDIR = ROOT / "out" / "figures"

TITLE = ("Presentation Format Determines Whether Language Models "
         "Respond to the Cost of Punishment")
AUTHOR = "Vladimir Vyatchanin — Letovo School, Moscow"

FIGURES = [
    ("fig1_slopes.png",
     "Figure 1.\tCost slope by source in branch A, with 95% confidence "
     "intervals. Humans are three to four times steeper than any model."),
    ("fig2_rates.png",
     "Figure 2.\tShare of punishment by price. (a) One condition per prompt; "
     "(b) all four conditions in a single prompt. The same models that look "
     "price-insensitive in (a) track price in (b)."),
    ("fig3_degenerate.png",
     "Figure 3.\tShare of cells in which every repetition gave the identical "
     "answer, by model."),
]


def drop_slide(prs: Presentation, index: int) -> None:
    """python-pptx has no delete; drop the id entry and its relationship."""
    slides = prs.slides._sldIdLst
    slide_id = list(slides)[index]
    prs.part.drop_rel(slide_id.rId)
    slides.remove(slide_id)


def main() -> int:
    missing = [f for f, _ in FIGURES if not (FIGDIR / f).exists()]
    if missing:
        raise SystemExit(f"missing figures: {missing}. Run make_figures.py first.")

    prs = Presentation(str(TEMPLATE))

    # Title slide: keep the layout, use our title.
    title_slide = prs.slides[0]
    placeholders = [s for s in title_slide.shapes if s.has_text_frame]
    if placeholders:
        placeholders[0].text_frame.text = TITLE
    if len(placeholders) > 1:
        placeholders[1].text_frame.text = AUTHOR

    figure_layout = prs.slides[1].slide_layout
    for _ in range(len(prs.slides._sldIdLst) - 1):
        drop_slide(prs, 1)

    slide_w = prs.slide_width
    for name, caption in FIGURES:
        slide = prs.slides.add_slide(figure_layout)

        # Caption goes in the layout's title placeholder, as in the template.
        caption_shape = slide.shapes.title
        caption_shape.text_frame.text = caption
        for para in caption_shape.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(12)

        top = Inches(1.45)
        max_h = prs.slide_height - top - Inches(0.25)
        picture = slide.shapes.add_picture(str(FIGDIR / name), 0, 0,
                                           width=Inches(8.2))
        if picture.height > max_h:  # tall figures must fit the slide, not spill
            scale = max_h / picture.height
            picture.height = int(picture.height * scale)
            picture.width = int(picture.width * scale)
        picture.left = int((slide_w - picture.width) / 2)
        picture.top = top

        # Anything the layout left empty would print as "Click to add text".
        for shape in list(slide.shapes):
            if (shape.has_text_frame and shape is not caption_shape
                    and not shape.text_frame.text.strip()):
                shape._element.getparent().remove(shape._element)

    prs.save(str(OUT))
    size_mb = OUT.stat().st_size / 1e6
    print(f"wrote {OUT}  ({size_mb:.2f} MB, {len(prs.slides._sldIdLst)} slides)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
